"""AIpediya Catalog Master: the canonical registry of every Model and Tool.

The workbook is the working source of truth. A record appears and is verified
here first; only PUBLISHED rows are later synchronised to Local/Production.
`import_from_local` brings Local records into the workbook without
overwriting master data; `validate` enforces the rules and reports drift.
Public Number is derived: chronological by verified release date among
PUBLISHED rows, recomputed on every import, with changes written to the
Changelog. Record ID (slug) never changes. Reads the database only.
"""
import json
import os
import re
import tempfile
import urllib.parse
import urllib.request
from datetime import date, datetime, timezone
from pathlib import Path

from .comparison import alphabet

WORKBOOK_PATH = Path("AI_CONTEXT") / "AIpediya_Model_Verification_Master.xlsx"
SCHEMA = "aipediya-catalog-master/4"
UPGRADABLE_SCHEMAS = {"aipediya-catalog-master/2", "aipediya-catalog-master/3"}
PRODUCTION_ORIGIN = "https://aipediya.com"
MAX_VALIDATED_ROW = 5000

STATUSES = ["PUBLISHED", "NEEDS_REVIEW"]
DECISIONS = ["PUBLIC", "ARCHIVE", "NEEDS_REVIEW"]
# Editorial reason codes per Publication Decision (column "Decision Code").
DECISION_CODES = {
    "PUBLIC": ["CURRENT_RELEASE", "HISTORICAL_RELEASE"],
    "ARCHIVE": [
        "DUPLICATE", "ALIAS_SNAPSHOT", "API_SNAPSHOT", "CONFIGURATION", "FORMAT_VARIANT",
        "TECHNICAL_CHECKPOINT", "FAMILY_AGGREGATE", "MODEL_APP_SPLIT", "OUT_OF_SCOPE",
        "CANCELLED", "TEST_ARTIFACT",
    ],
    "NEEDS_REVIEW": ["IDENTITY", "DATE", "EXISTENCE", "SCOPE", "SOURCE"],
}
ALL_DECISION_CODES = [code for codes in DECISION_CODES.values() for code in codes]
# "Canonical / Parent Record ID" semantics. SAME_ENTITY relations mean the row
# is the same thing as its canonical card under another name/form: its aliases
# may lead (redirect) to that card. The others only relate distinct entities.
RELATION_TYPES = ["DUPLICATE_OF", "ALIAS_OF", "FORMAT_OF", "MODE_OF", "SNAPSHOT_OF", "VARIANT_OF", "MEMBER_OF"]
SAME_ENTITY = {"DUPLICATE_OF", "ALIAS_OF", "FORMAT_OF", "MODE_OF"}
CODES_NEEDING_CANONICAL = {"DUPLICATE", "ALIAS_SNAPSHOT", "FORMAT_VARIANT", "CONFIGURATION"}
ALIAS_SEPARATOR = " | "
# An empty cell means "no new information"; this token explicitly clears a
# value when the master is synchronised to Local.
CLEAR = "<CLEAR>"
MAX_RECORD_ID = 50  # Django SlugField length used by Local
OFFER_UNITS = [
    "input", "output", "image", "megapixel", "second", "minute", "month", "year", "hour",
    "million_characters", "thousand_characters", "other", "request", "credit", "cache_read", "cache_write",
]
PRECISIONS = ["day", "month", "year"]
YES_NO = ["YES", "NO"]
MODEL_CATEGORIES = ["text", "image", "video", "audio", "other"]
TOOL_CATEGORIES = [
    "ai_app", "coding_assistant", "coding_agent", "runtime", "api_platform",
    "api_service", "ide_tool", "client", "agent_platform", "creative_app",
]
CATALOG_STATUSES = ["active", "deprecated", "retired", "archived"]
STAGES = ["released", "preview", "beta", "research"]
LOCAL_EXECUTION = ["yes", "no", "hybrid"]

WORKFLOW = [
    "Record ID", "Status", "Publication Decision", "Public Number", "Name", "Developer",
    "Exact Release Date", "Approx Date", "Approx Precision",
]
VERIFICATION = [
    "Missing Data", "Reason", "Decision Code", "Decision Date", "Decision Sources",
    "Canonical / Parent Record ID", "Relation Type", "Aliases", "Official Source",
    "Secondary Source", "Last Verified", "On Local", "On Production", "Notes",
]
MODEL_PUBLIC = [
    "Release Stage", "Family", "Version", "Category", "Tasks", "Input Modalities",
    "Output Modalities", "Context", "License", "Open Weights", "Catalog Status",
    "Developer Country", "Origin Countries", "Description EN", "Description RU",
    "Suitable EN", "Suitable RU", "Limitations EN", "Limitations RU",
    "Source Title", "Source URL", "Source Publisher", "Checked (DB)",
    "Release Evidence (JSON)", "Approx Evidence (JSON)", "Origin (JSON)",
    "Philosophy (JSON)", "Description Other Locales (JSON)",
    "Suitable Other Locales (JSON)", "Limitations Other Locales (JSON)",
    "Research Entity ID",
]
TOOL_PUBLIC = [
    "Version", "Category", "Purposes", "Local Execution", "Official URL",
    "Platforms", "Supported Models", "Catalog Status", "Developer Country",
    "Description EN", "Description RU", "Ecosystem EN", "Ecosystem RU",
    "Source Title", "Source URL", "Source Publisher", "Checked (DB)",
    "Release Evidence (JSON)", "Approx Evidence (JSON)",
    "Description Other Locales (JSON)", "Ecosystem Other Locales (JSON)",
    "Legacy Record ID",
]
MODEL_SERVICE = [
    "Local Number", "Local PK", "Offers", "Evaluations", "Access",
    "Research Date Hint", "Import Sources (unverified)", "Import Batch", "Missing Count",
]
TOOL_SERVICE = [
    "Local Number", "Local PK", "Offers", "Access",
    "Research Date Hint", "Import Sources (unverified)", "Import Batch", "Missing Count",
]
MAIN = {
    "Models": WORKFLOW + VERIFICATION + MODEL_PUBLIC + MODEL_SERVICE,
    "Tools": WORKFLOW + VERIFICATION + TOOL_PUBLIC + TOOL_SERVICE,
}
# Public-data columns compared against Local to report pending sync (drift).
PUBLIC = {
    "Models": ["Name", "Developer", "Exact Release Date", "Approx Date"] + MODEL_PUBLIC,
    "Tools": ["Name", "Developer", "Exact Release Date", "Approx Date"] + TOOL_PUBLIC,
}
LOCAL_OWNED = {
    "On Local", "Local Number", "Local PK", "Research Date Hint",
    "Import Sources (unverified)", "Import Batch",
}
AUX = {
    "Offers": [
        "Key", "Record Type", "Record ID", "Service", "Service Kind", "Service URL",
        "Provider", "Compute Location", "Amount", "Unit", "Billing Unit",
        "Conditions EN", "Conditions RU", "Conditions Extra (JSON)", "Primary",
        "Active", "Source URL", "Checked", "Research Key",
    ],
    "Evaluations": [
        "Key", "Record Type", "Record ID", "Benchmark", "Protocol",
        "Benchmark Category", "Unit", "Higher Is Better", "Score", "Evaluator",
        "Result Kind", "Independent", "Public", "Measured", "Configuration",
        "Confidence Low", "Confidence High", "Conditions EN", "Conditions RU",
        "Conditions Extra (JSON)", "Source URL", "Checked", "Observation Key",
        "Source Model", "Source Record ID", "Snapshot", "Source SHA256",
    ],
    "Access": [
        "Key", "Record Type", "Record ID", "Service", "Service Kind", "Service URL",
        "Provider", "Compute Location", "Source URL", "Checked",
    ],
    "Facts": ["Key", "Record Type", "Record ID", "Fact", "Value (JSON)", "Source URL", "Checked"],
    "Origins": ["Key", "Record Type", "Record ID", "Country", "Position", "Source URL", "Checked"],
    "Tool Platforms": ["Key", "Record Type", "Record ID", "Platform", "Source URL", "Checked"],
}
CHANGELOG = ["Timestamp (UTC)", "Sheet", "Record ID", "Field", "Before", "After", "Reason"]
RECORD_TYPE_SHEET = {"model": "Models", "tool": "Tools"}

RECORD_ID = re.compile(r"^[-a-zA-Z0-9_]+$")
ISO_DATE = re.compile(r"^\d{4}-\d{2}-\d{2}$")
APPROX_DATE = re.compile(r"^≈\s?\d{4}(-\d{2}(-\d{2})?)?$")


def _text(value):
    if value is None:
        return ""
    if isinstance(value, bool):
        return "YES" if value else "NO"
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _json(value):
    return json.dumps(value, ensure_ascii=False, sort_keys=True) if value else ""


def _locales(value):
    value = value or {}
    other = {k: v for k, v in value.items() if k not in ("en", "ru")}
    return _text(value.get("en")), _text(value.get("ru")), _json(other)


def _decimal(value):
    if value is None:
        return ""
    text = format(value.normalize(), "f")
    return text


def _approx(released, precision):
    if not released:
        return ""
    return "≈" + (released.strftime("%Y-%m") if precision == "month" else released.isoformat())


def _valid_iso(value):
    if not ISO_DATE.match(value):
        return False
    try:
        date.fromisoformat(value)
    except ValueError:
        return False
    return True


def now_utc():
    return datetime.now(timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ")


# ---------------------------------------------------------------- Local rows

def _research_columns(obj):
    evidence = obj.release_evidence or {}
    hint = evidence.get("date_text") or ""
    if hint and evidence.get("precision"):
        hint += " (%s)" % evidence["precision"]
    sources = []
    for url in [evidence.get("source_url"), *(evidence.get("checked_sources") or []),
                obj.source.url if obj.source_id else ""]:
        if url and url not in sources:
            sources.append(url)
    return {
        "Research Date Hint": hint,
        "Import Sources (unverified)": " | ".join(sources),
        "Import Batch": evidence.get("import_batch", ""),
    }


def _verification(obj):
    """Verification metadata for a record first imported from Local.

    Published records carry the audited evidence already used by the site;
    hidden research rows get none, because their imported links are unverified.
    """
    evidence = obj.release_evidence or {}
    if not obj.published:
        notes = ["Hidden in Local (research layer); existence, date and sources are not verified."]
        if evidence.get("reason"):
            notes.append(evidence["reason"])
        audit = evidence.get("audit_notes") or {}
        notes.extend(t for t in (audit.get("text") or []) if t)
        if audit.get("media"):
            notes.append(audit["media"])
        return {"Status": "NEEDS_REVIEW", "Publication Decision": "NEEDS_REVIEW",
                "Decision Code": "SOURCE", "Decision Date": now_utc()[:10],
                "Reason": "Скрыт в Local (исследовательский слой): существование, дата и источники не проверены.",
                "Official Source": "", "Secondary Source": "",
                "Last Verified": "", "Notes": " ".join(notes)}
    official = evidence.get("source_url") or (obj.source.url if obj.source_id else "")
    secondary = obj.source.url if obj.source_id and obj.source.url != official else ""
    return {"Status": "PUBLISHED", "Publication Decision": "PUBLIC",
            "Official Source": official, "Secondary Source": secondary,
            "Last Verified": evidence.get("checked") or _text(obj.checked), "Notes": ""}


def model_row(model):
    description, description_ru, description_other = _locales(model.description)
    suitable, suitable_ru, suitable_other = _locales(model.suitable)
    limitations, limitations_ru, limitations_other = _locales(model.limitations)
    developer = model.family.developer
    row = {
        "Record ID": model.slug, "Name": model.name, "Developer": developer.name,
        "Exact Release Date": _text(model.released),
        "Approx Date": _approx(model.approx_released, model.approx_precision),
        "On Local": _text(model.published),
        "Release Stage": model.release_stage, "Family": model.family.name,
        "Version": model.version, "Category": model.category,
        "Tasks": "; ".join(model.tasks or []),
        "Input Modalities": "; ".join(model.input_modalities or []),
        "Output Modalities": "; ".join(model.output_modalities or []),
        "Context": _text(model.context), "License": model.license,
        "Open Weights": _text(model.open_weights), "Catalog Status": model.catalog_status,
        "Developer Country": developer.country,
        "Origin Countries": "; ".join(link.country_id for link in model.origin_country_links.all()),
        "Description EN": description, "Description RU": description_ru,
        "Suitable EN": suitable, "Suitable RU": suitable_ru,
        "Limitations EN": limitations, "Limitations RU": limitations_ru,
        "Source Title": model.source.title, "Source URL": model.source.url,
        "Source Publisher": model.source.publisher, "Checked (DB)": _text(model.checked),
        "Release Evidence (JSON)": _json(model.release_evidence),
        "Approx Evidence (JSON)": _json(model.approx_evidence),
        "Origin (JSON)": _json(model.origin), "Philosophy (JSON)": _json(model.philosophy),
        "Description Other Locales (JSON)": description_other,
        "Suitable Other Locales (JSON)": suitable_other,
        "Limitations Other Locales (JSON)": limitations_other,
        "Research Entity ID": model.research_entity_id or "",
        "Local Number": _text(model.public_number), "Local PK": _text(model.pk),
    }
    row.update(_research_columns(model))
    return row


def tool_row(tool):
    description, description_ru, description_other = _locales(tool.description)
    ecosystem, ecosystem_ru, ecosystem_other = _locales(tool.ecosystem)
    row = {
        "Record ID": tool.slug, "Name": tool.name, "Developer": tool.developer.name,
        "Exact Release Date": _text(tool.released),
        "Approx Date": _approx(tool.approx_released, tool.approx_precision),
        "On Local": _text(tool.published),
        "Version": tool.version, "Category": tool.category,
        "Purposes": "; ".join(tool.purposes or []), "Local Execution": tool.local_execution,
        "Official URL": tool.official_url,
        "Platforms": "; ".join(link.platform.code for link in tool.platform_links.all()),
        "Supported Models": "; ".join(link.model.slug for link in tool.model_links.all()),
        "Catalog Status": tool.catalog_status, "Developer Country": tool.developer.country,
        "Description EN": description, "Description RU": description_ru,
        "Ecosystem EN": ecosystem, "Ecosystem RU": ecosystem_ru,
        "Source Title": tool.source.title, "Source URL": tool.source.url,
        "Source Publisher": tool.source.publisher, "Checked (DB)": _text(tool.checked),
        "Release Evidence (JSON)": _json(tool.release_evidence),
        "Approx Evidence (JSON)": _json(tool.approx_evidence),
        "Description Other Locales (JSON)": description_other,
        "Ecosystem Other Locales (JSON)": ecosystem_other,
        "Legacy Record ID": tool.legacy_version.slug if tool.legacy_version_id else "",
        "Local Number": _text(tool.public_number), "Local PK": _text(tool.pk),
    }
    row.update(_research_columns(tool))
    return row


def _owner(model_version, tool_by_legacy):
    """(record type, record id) for a ModelVersion-owned relation."""
    tool = tool_by_legacy.get(model_version.pk)
    if tool is not None:
        return "tool", tool.slug
    return "model", model_version.slug


def _service(service):
    return {
        "Service": service.name, "Service Kind": service.kind, "Service URL": service.url,
        "Provider": service.provider.name, "Compute Location": service.compute_location,
    }


def local_snapshot():
    """Every Local Model/Tool record and relation as master rows."""
    from .models import (
        Access, Evaluation, Fact, ModelOriginCountry, ModelVersion, Offer, Tool, ToolPlatform,
    )

    tools = list(Tool.objects.select_related("developer", "source", "legacy_version")
                 .prefetch_related("platform_links__platform", "model_links__model")
                 .order_by("public_number", "slug"))
    tool_by_legacy = {t.legacy_version_id: t for t in tools if t.legacy_version_id}
    models = list(ModelVersion.objects.filter(entry_type="model")
                  .select_related("family__developer", "source")
                  .prefetch_related("origin_country_links")
                  .order_by("public_number", "slug"))
    snapshot = {
        "Models": {m.slug: (model_row(m), _verification(m)) for m in models},
        "Tools": {t.slug: (tool_row(t), _verification(t)) for t in tools},
    }
    aux = {name: {} for name in AUX}
    for offer in Offer.objects.select_related("model", "service__provider", "source").order_by("pk"):
        kind, record = _owner(offer.model, tool_by_legacy)
        en, ru, extra = _locales(offer.conditions)
        aux["Offers"]["offer-%d" % offer.pk] = {
            "Record Type": kind, "Record ID": record, **_service(offer.service),
            "Amount": _decimal(offer.amount), "Unit": offer.unit, "Billing Unit": offer.billing_unit,
            "Conditions EN": en, "Conditions RU": ru, "Conditions Extra (JSON)": extra,
            "Primary": _text(offer.primary), "Active": _text(offer.active),
            "Source URL": offer.source.url, "Checked": _text(offer.checked),
            "Research Key": offer.research_key or "",
        }
    for ev in Evaluation.objects.select_related("model", "benchmark", "source").order_by("pk"):
        kind, record = _owner(ev.model, tool_by_legacy)
        en, ru, extra = _locales(ev.conditions)
        aux["Evaluations"]["evaluation-%d" % ev.pk] = {
            "Record Type": kind, "Record ID": record, "Benchmark": ev.benchmark.name,
            "Protocol": ev.benchmark.protocol, "Benchmark Category": ev.benchmark.category,
            "Unit": ev.benchmark.unit, "Higher Is Better": _text(ev.benchmark.higher_is_better),
            "Score": _decimal(ev.score), "Evaluator": ev.evaluator, "Result Kind": ev.result_kind,
            "Independent": _text(ev.independent), "Public": _text(ev.public),
            "Measured": _text(ev.measured), "Configuration": ev.configuration,
            "Confidence Low": _decimal(ev.confidence_low), "Confidence High": _decimal(ev.confidence_high),
            "Conditions EN": en, "Conditions RU": ru, "Conditions Extra (JSON)": extra,
            "Source URL": ev.source.url, "Checked": _text(ev.checked),
            "Observation Key": ev.observation_key or "", "Source Model": ev.source_model,
            "Source Record ID": ev.source_record_id, "Snapshot": ev.snapshot,
            "Source SHA256": ev.source_sha256,
        }
    for access in Access.objects.select_related("model", "service__provider", "source").order_by("pk"):
        kind, record = _owner(access.model, tool_by_legacy)
        aux["Access"]["access-%d" % access.pk] = {
            "Record Type": kind, "Record ID": record, **_service(access.service),
            "Source URL": access.source.url, "Checked": _text(access.checked),
        }
    for fact in Fact.objects.select_related("model", "source").order_by("pk"):
        kind, record = _owner(fact.model, tool_by_legacy)
        aux["Facts"]["fact-%d" % fact.pk] = {
            "Record Type": kind, "Record ID": record, "Fact": fact.key,
            "Value (JSON)": _json(fact.value), "Source URL": fact.source.url,
            "Checked": _text(fact.checked),
        }
    for link in ModelOriginCountry.objects.select_related("model", "source").order_by("pk"):
        aux["Origins"]["origin-%d" % link.pk] = {
            "Record Type": "model", "Record ID": link.model.slug, "Country": link.country_id,
            "Position": _text(link.position), "Source URL": link.source.url,
            "Checked": _text(link.checked),
        }
    for link in ToolPlatform.objects.select_related("tool", "platform", "source").order_by("pk"):
        aux["Tool Platforms"]["platform-%d" % link.pk] = {
            "Record Type": "tool", "Record ID": link.tool.slug, "Platform": link.platform.code,
            "Source URL": link.source.url, "Checked": _text(link.checked),
        }
    return snapshot, aux


def production_sitemaps(get):
    """Sitemap documents to scan: /sitemap.xml itself, or, when it is a
    sitemap index (locale paths since GSD-1.0), its English root sitemap —
    every locale lists the same public entities."""
    root = get("/sitemap.xml")
    if "<sitemapindex" not in root:
        return [root]
    children = re.findall(r"<loc>([^<]+)</loc>", root)
    english = [u for u in children if u.rstrip("/").endswith("/en.xml")] or children[:1]
    return [get(urllib.parse.urlsplit(url).path) for url in english]


def fetch_production(origin=PRODUCTION_ORIGIN, timeout=60):
    """Read-only: slugs listed in the public sitemap plus the live release."""
    def get(path):
        request = urllib.request.Request(origin + path, headers={"User-Agent": "aipedia-catalog-master"})
        with urllib.request.urlopen(request, timeout=timeout) as response:
            return response.read().decode("utf-8")

    found = {"Models": set(), "Tools": set()}
    for sitemap in production_sitemaps(get):
        for kind, slug in re.findall(r"<loc>[^<]*?/(models|tools)/([^?<]+)", sitemap):
            found["Models" if kind == "models" else "Tools"].add(slug)
    release = ""
    try:
        release = json.loads(get("/healthz")).get("release", "")
    except (OSError, ValueError):
        pass
    return found, release


# ------------------------------------------------------------ master logic

def missing_items(sheet, row):
    row = {k: ("" if v == CLEAR else v) for k, v in row.items()}
    if row.get("Publication Decision") == "ARCHIVE":
        # Archived rows need a justified, dated, sourced decision (and their
        # canonical card when they are a duplicate/alias), not publication data.
        items = [key for column, key in (
            ("Decision Code", "decision_code"), ("Reason", "reason"),
            ("Decision Date", "decision_date"), ("Last Verified", "last_verified")) if not row.get(column)]
        if not row.get("Official Source") and not row.get("Decision Sources"):
            items.append("decision_sources")
        if row.get("Decision Code") in CODES_NEEDING_CANONICAL and not row.get("Canonical / Parent Record ID"):
            items.append("canonical_record")
        return items
    items = []
    if not row.get("Publication Decision"):
        items.append("publication_decision")
    for column, key in (("Name", "name"), ("Developer", "developer"), ("Category", "category")):
        if not row.get(column):
            items.append(key)
    if not row.get("Exact Release Date") and not row.get("Approx Date"):
        items.append("date (exact or ≈)")
    if not row.get("Official Source"):
        items.append("official_source")
    if not row.get("Last Verified"):
        items.append("last_verified")
    if not row.get("Description EN"):
        items.append("description_en")
    return items


def split_aliases(value):
    """Aliases cell -> list of names; the separator is ' | ' (pipe)."""
    return [part.strip() for part in (value or "").split("|") if part.strip()]


def _alias_key(name):
    return " ".join((name or "").casefold().split())


def alias_index(rows):
    """Casefolded name/alias -> set of Record IDs claiming it (one sheet)."""
    index = {}
    for row in rows:
        for name in [row.get("Name", "")] + split_aliases(row.get("Aliases")):
            if name:
                index.setdefault(_alias_key(name), set()).add(row["Record ID"])
    return index


def cross_parent(value):
    """'Models:<id>' / 'Tools:<id>' -> (sheet, id) for a canonical card on the
    other sheet (e.g. a model wrongly listed as a tool), else None."""
    sheet, sep, record = (value or "").partition(":")
    return (sheet, record) if sep and sheet in MAIN and record else None


def canonical_target(row, by_id):
    """Record ID a SAME_ENTITY row resolves to (following the chain), or ''."""
    seen, current = set(), row
    while current and current.get("Relation Type") in SAME_ENTITY and current.get("Canonical / Parent Record ID"):
        if current["Record ID"] in seen:
            return ""
        seen.add(current["Record ID"])
        current = by_id.get(current["Canonical / Parent Record ID"])
    return current["Record ID"] if current and current is not row else ""


def parent_cycles(rows):
    """Record IDs that sit on a Canonical / Parent cycle."""
    parent = {r["Record ID"]: r.get("Canonical / Parent Record ID", "") for r in rows if r.get("Record ID")}
    cyclic = set()
    for start in parent:
        path, node = [], start
        while node and node in parent and node not in path:
            path.append(node)
            node = parent[node]
        if node in path:
            cyclic.update(path[path.index(node):])
    return cyclic


def chronology_date(row):
    value = row.get("Exact Release Date") or row.get("Approx Date", "").lstrip("≈").strip()
    if not value:
        return None
    parts = value.split("-")
    try:
        return date(int(parts[0]), int(parts[1]) if len(parts) > 1 else 1,
                    int(parts[2]) if len(parts) > 2 else 1)
    except ValueError:
        return None


def approx_precision(value):
    value = (value or "").lstrip("≈").strip()
    if not value:
        return ""
    return {1: "year", 2: "month", 3: "day"}.get(len(value.split("-")), "")


def chronology_numbers(rows):
    """Record ID -> Public Number for PUBLISHED rows with a verified date."""
    dated = [r for r in rows if r.get("Status") == "PUBLISHED"
             and r.get("Publication Decision") == "PUBLIC" and chronology_date(r)]
    dated.sort(key=lambda r: (chronology_date(r), alphabet(r.get("Name", "")), r["Record ID"]))
    return {r["Record ID"]: str(n) for n, r in enumerate(dated, 1)}


def new_record_id(sheet, name, taken):
    """Controlled Record ID for a new row: slug of the name + stable hash."""
    import hashlib
    base = re.sub(r"-+", "-", re.sub(r"[^a-z0-9]+", "-", (name or "").lower())).strip("-")[:MAX_RECORD_ID - 9]
    base = base.strip("-") or "record"
    salt = 0
    while True:
        digest = hashlib.sha1(("aipediya-master:%s:%s:%d" % (sheet, name, salt)).encode()).hexdigest()[:8]
        record = "%s-%s" % (base, digest)
        if record not in taken:
            return record
        salt += 1


def normalize_drafts(workbook_rows, changelog, stamp):
    """Rows typed by a person with only some known facts become drafts.

    A row with a Name but no Record ID gets a controlled Record ID (kept from
    then on); empty Status / Publication Decision become NEEDS_REVIEW with a
    draft note, so an incomplete row is saved but never counts as publishable.
    """
    touched = 0
    for sheet in MAIN:
        rows = workbook_rows.get(sheet, [])
        taken = {r.get("Record ID") for r in rows if r.get("Record ID")}
        for row in rows:
            changes = {}
            if not row.get("Record ID"):
                if not row.get("Name"):
                    continue
                changes["Record ID"] = new_record_id(sheet, row["Name"], taken)
                taken.add(changes["Record ID"])
            if not row.get("Status"):
                changes["Status"] = "NEEDS_REVIEW"
            if not row.get("Publication Decision"):
                changes["Publication Decision"] = "NEEDS_REVIEW"
                if not row.get("Decision Code"):
                    changes["Decision Code"] = "SOURCE" if not row.get("Official Source") else "IDENTITY"
                if not row.get("Reason"):
                    changes["Reason"] = ("Черновик: строка добавлена вручную %s; нужна проверка идентичности, "
                                         "источника и даты." % stamp[:10])
                if not row.get("Decision Date"):
                    changes["Decision Date"] = stamp[:10]
            record = changes.get("Record ID") or row["Record ID"]
            for field, value in changes.items():
                changelog.append({"Timestamp (UTC)": stamp, "Sheet": sheet, "Record ID": record, "Field": field,
                                  "Before": row.get(field, ""), "After": value, "Reason": "draft row normalised"})
                row[field] = value
            touched += bool(changes)
    return touched


def refresh_derived(workbook_rows, changelog, reason, stamp, quiet_ids=()):
    """Recompute Missing Data, counts and Public Numbers; log number changes."""
    normalize_drafts(workbook_rows, changelog, stamp)
    counts = {}
    for aux_sheet in ("Offers", "Evaluations", "Access"):
        for row in workbook_rows.get(aux_sheet, []):
            key = (RECORD_TYPE_SHEET.get(row.get("Record Type"), ""), row.get("Record ID"), aux_sheet)
            counts[key] = counts.get(key, 0) + 1
    changed = 0
    for sheet in MAIN:
        rows = workbook_rows[sheet]
        numbers = chronology_numbers(rows)
        for row in rows:
            row["Approx Precision"] = approx_precision(row.get("Approx Date"))
            items = missing_items(sheet, row)
            row["Missing Data"] = "; ".join(items)
            row["Missing Count"] = str(len(items))
            for aux_sheet in ("Offers", "Evaluations", "Access"):
                if aux_sheet in MAIN[sheet]:
                    row[aux_sheet] = str(counts.get((sheet, row["Record ID"], aux_sheet), 0))
            new = numbers.get(row["Record ID"], "")
            if row.get("Public Number", "") != new:
                if row["Record ID"] not in quiet_ids:
                    changelog.append({
                        "Timestamp (UTC)": stamp, "Sheet": sheet, "Record ID": row["Record ID"],
                        "Field": "Public Number", "Before": row.get("Public Number", ""),
                        "After": new, "Reason": reason,
                    })
                    changed += 1
                row["Public Number"] = new
    return changed


# Meta labels once filled with finite-range Excel formulas (v004-v012). They are
# now computed values, refreshed on every write, so added rows are never missed.
COMPUTED_META = {
    "Записей в каталоге": "records",
    "Tools: охвачено проверкой": "covered_tools",
    "Models: охвачено проверкой": "covered_models",
    "Всего охвачено проверкой": "covered",
    "Охват, не полнота проверки": "coverage_percent",
    "Ещё не проверялись": "not_covered",
    "Всего Tools": "tools",
    "Tools с датой": "tools_dated",
    "Tools без даты": "tools_undated",
    "Заполненность дат Tools": "tools_dated_percent",
    "Models: обязательные пробелы": "models_gaps",
    "Предложено в архив, всего": "archive",
}
OBSOLETE_META = {"Без даты: ещё в очереди", "Без даты: проверено, не решено",
                 "Проверено в этом проходе", "Проверено ранее"}


def catalog_stats(workbook_rows):
    """Reproducible counters; coverage counts unique (Sheet, Record ID) pairs."""
    ids = {sheet: {r["Record ID"] for r in workbook_rows.get(sheet, [])} for sheet in MAIN}
    covered = {sheet: set() for sheet in MAIN}
    for event in workbook_rows.get("Changelog", []):
        if event.get("Field") == "Verification result" and event.get("Record ID") in ids.get(event.get("Sheet"), ()):
            covered[event["Sheet"]].add(event["Record ID"])
    total = sum(len(v) for v in ids.values())
    done = sum(len(v) for v in covered.values())
    tools = workbook_rows.get("Tools", [])
    dated = sum(1 for r in tools if chronology_date(r))
    stats = {
        "records": total, "covered_models": len(covered["Models"]), "covered_tools": len(covered["Tools"]),
        "covered": done, "coverage_percent": "%.2f%%" % (100.0 * done / total if total else 0),
        "not_covered": total - done, "tools": len(tools), "tools_dated": dated,
        "tools_undated": len(tools) - dated,
        "tools_dated_percent": "%.2f%%" % (100.0 * dated / len(tools) if tools else 0),
        "models_gaps": sum(1 for r in workbook_rows.get("Models", []) if r.get("Missing Data")),
        "archive": sum(1 for s in MAIN for r in workbook_rows.get(s, []) if r.get("Publication Decision") == "ARCHIVE"),
    }
    for sheet in MAIN:
        for decision in DECISIONS:
            stats["%s %s" % (sheet, decision)] = sum(
                1 for r in workbook_rows.get(sheet, []) if r.get("Publication Decision") == decision)
        stats["%s PUBLISHED" % sheet] = sum(1 for r in workbook_rows.get(sheet, []) if r.get("Status") == "PUBLISHED")
    return stats


def refresh_meta(meta, workbook_rows, stamp):
    """Replace formula/obsolete Meta rows by computed values; add a summary."""
    stats = catalog_stats(workbook_rows)
    refreshed = {}
    for key, value in meta.items():
        if key in OBSOLETE_META:
            continue
        if key in COMPUTED_META:
            value = str(stats[COMPUTED_META[key]])
        elif str(value).startswith("="):
            continue
        refreshed[key] = value
    refreshed["Computed (UTC)"] = stamp
    refreshed["Computed counts"] = "; ".join(
        "%s=%s" % (k, stats[k]) for k in stats if k.split(" ")[0] in MAIN)
    refreshed["Coverage rule"] = ("Охват = уникальные пары (Sheet, Record ID) с событием Verification result "
                                  "в Changelog; повторная проверка не увеличивает процент; это не полнота полей.")
    return refreshed


# Link rows identified by content: a link already in the master under a
# master-made key is not imported again under its Local key.
NATURAL_KEYS = {"Tool Platforms": ("Record ID", "Platform"), "Origins": ("Record ID", "Country")}


def import_from_local(workbook_rows, snapshot, aux, production=None, stamp=None):
    """Merge Local into the master; master data wins for existing records.

    Returns a summary dict. `production` is (found slugs per sheet, release)
    from fetch_production, or None to keep existing On Production values.
    """
    stamp = stamp or now_utc()
    changelog = workbook_rows.setdefault("Changelog", [])
    added_ids, summary = set(), {}
    for sheet, columns in MAIN.items():
        rows = workbook_rows.setdefault(sheet, [])
        known = {r["Record ID"] for r in rows}
        for row in rows:
            local = snapshot[sheet].get(row["Record ID"])
            if local is not None:
                row.update({k: v for k, v in local[0].items() if k in LOCAL_OWNED})
            elif row.get("Record ID"):
                row["On Local"] = "NO"
        added = 0
        for record_id, (local_row, verification) in snapshot[sheet].items():
            if record_id in known:
                continue
            row = {c: "" for c in columns}
            row.update(local_row)
            row.update(verification)
            rows.append(row)
            added_ids.add(record_id)
            added += 1
        if production is not None:
            found = production[0][sheet]
            for row in rows:
                row["On Production"] = "YES" if row["Record ID"] in found else "NO"
        if added:
            changelog.append({
                "Timestamp (UTC)": stamp, "Sheet": sheet, "Record ID": "*", "Field": "rows",
                "Before": str(len(rows) - added), "After": str(len(rows)),
                "Reason": "import from Local: added %d records" % added,
            })
        summary[sheet] = {"rows": len(rows), "added": added}
    for sheet in AUX:
        rows = workbook_rows.setdefault(sheet, [])
        known = {r["Key"] for r in rows}
        natural = NATURAL_KEYS.get(sheet)
        if natural:
            known_natural = {tuple(r.get(c, "") for c in natural) for r in rows}
        new = [{"Key": key, **values} for key, values in aux[sheet].items() if key not in known
               and not (natural and tuple(values.get(c, "") for c in natural) in known_natural)]
        rows.extend(new)
        if new:
            changelog.append({
                "Timestamp (UTC)": stamp, "Sheet": sheet, "Record ID": "*", "Field": "rows",
                "Before": str(len(rows) - len(new)), "After": str(len(rows)),
                "Reason": "import from Local: added %d rows" % len(new),
            })
        summary[sheet] = {"rows": len(rows), "added": len(new)}
    summary["renumbered"] = refresh_derived(
        workbook_rows, changelog, "chronology recomputed on import", stamp, quiet_ids=added_ids)
    return summary


def validate(workbook_rows, snapshot=None, aux=None):
    """Return (errors, warnings, drift) for the master workbook."""
    errors, warnings, drift = [], [], {}
    ids = {}
    cross_parents = []
    for sheet in MAIN:
        categories = MODEL_CATEGORIES if sheet == "Models" else TOOL_CATEGORIES
        rows = workbook_rows.get(sheet, [])
        numbers = chronology_numbers(rows)
        seen, parents = {}, []
        for index, row in enumerate(rows, start=2):
            row = {k: ("" if v == CLEAR else v) for k, v in row.items()}
            record = row.get("Record ID", "")
            where = "%s row %d (%s)" % (sheet, index, record or row.get("Name") or "?")
            if not record:
                errors.append("%s: empty Record ID" % where)
                continue
            if len(record) > MAX_RECORD_ID:
                errors.append("%s: Record ID longer than %d characters" % (where, MAX_RECORD_ID))
            if not RECORD_ID.match(record):
                errors.append("%s: Record ID may contain only letters, digits, '-' and '_'" % where)
            if record in seen:
                errors.append("%s: duplicate Record ID (row %d)" % (where, seen[record]))
            seen[record] = index
            status = row.get("Status", "")
            if status not in STATUSES:
                errors.append("%s: Status must be PUBLISHED or NEEDS_REVIEW" % where)
            exact, approx = row.get("Exact Release Date", ""), row.get("Approx Date", "")
            if exact and not _valid_iso(exact):
                errors.append("%s: Exact Release Date must be YYYY-MM-DD" % where)
            if approx and not APPROX_DATE.match(approx):
                errors.append("%s: Approx Date must look like ≈YYYY, ≈YYYY-MM or ≈YYYY-MM-DD" % where)
            if exact and approx:
                errors.append("%s: fill either Exact Release Date or Approx Date, not both" % where)
            if row.get("Last Verified") and not _valid_iso(row["Last Verified"]):
                errors.append("%s: Last Verified must be YYYY-MM-DD" % where)
            for column in ("On Local", "On Production", "Open Weights"):
                if row.get(column) and row[column] not in YES_NO:
                    errors.append("%s: %s must be YES/NO" % (where, column))
            if row.get("Category") and row["Category"] not in categories:
                errors.append("%s: unknown Category %s" % (where, row["Category"]))
            if row.get("Catalog Status") and row["Catalog Status"] not in CATALOG_STATUSES:
                errors.append("%s: unknown Catalog Status %s" % (where, row["Catalog Status"]))
            if row.get("Release Stage") and row["Release Stage"] not in STAGES:
                errors.append("%s: unknown Release Stage %s" % (where, row["Release Stage"]))
            if row.get("Local Execution") and row["Local Execution"] not in LOCAL_EXECUTION:
                errors.append("%s: unknown Local Execution %s" % (where, row["Local Execution"]))
            decision = row.get("Publication Decision", "")
            if decision not in DECISIONS:
                errors.append("%s: Publication Decision must be PUBLIC, ARCHIVE or NEEDS_REVIEW" % where)
            if status == "PUBLISHED" and decision != "PUBLIC":
                errors.append("%s: PUBLISHED requires Publication Decision PUBLIC" % where)
            code = row.get("Decision Code", "")
            if code and code not in DECISION_CODES.get(decision, []):
                errors.append("%s: Decision Code %s is not valid for %s" % (where, code, decision or "?"))
            if row.get("Decision Date") and not _valid_iso(row["Decision Date"]):
                errors.append("%s: Decision Date must be YYYY-MM-DD" % where)
            if decision == "ARCHIVE":
                lacking = [c for c in ("Reason", "Decision Code", "Decision Date") if not row.get(c)]
                if not row.get("Official Source") and not row.get("Decision Sources"):
                    lacking.append("Official Source or Decision Sources")
                if code in CODES_NEEDING_CANONICAL and not row.get("Canonical / Parent Record ID"):
                    lacking.append("Canonical / Parent Record ID")
                if lacking:
                    errors.append("%s: ARCHIVE requires %s" % (where, ", ".join(lacking)))
            elif decision == "NEEDS_REVIEW" and not row.get("Reason"):
                errors.append("%s: NEEDS_REVIEW requires the open question in Reason" % where)
            relation = row.get("Relation Type", "")
            if relation and relation not in RELATION_TYPES:
                errors.append("%s: unknown Relation Type %s" % (where, relation))
            parent = row.get("Canonical / Parent Record ID", "")
            if relation and not parent:
                errors.append("%s: Relation Type without Canonical / Parent Record ID" % where)
            if relation in SAME_ENTITY and decision != "ARCHIVE":
                errors.append("%s: %s rows are the same entity as their canonical card and must be ARCHIVE" % (where, relation))
            if parent == record:
                errors.append("%s: Canonical / Parent Record ID points to itself" % where)
            elif cross_parent(parent):
                cross_parents.append((sheet, where, record, cross_parent(parent), relation, decision))
            elif parent:
                parents.append((where, parent, relation))
            if status == "PUBLISHED":
                lacking = [c for c in ("Name", "Developer", "Official Source", "Last Verified") if not row.get(c)]
                if lacking:
                    errors.append("%s: PUBLISHED requires %s" % (where, ", ".join(lacking)))
                if not chronology_date(row):
                    warnings.append(("PUBLISHED without verified date: no Public Number", where))
                if row.get("On Local") == "NO":
                    warnings.append(("PUBLISHED but not public on Local: pending sync", where))
                if row.get("On Production") == "NO":
                    warnings.append(("PUBLISHED but not on Production: pending release", where))
            elif status == "NEEDS_REVIEW":
                for column, place in (("On Local", "Local"), ("On Production", "Production")):
                    if row.get(column) == "YES":
                        warnings.append(("not PUBLISHED (%s) but public on %s: hide at next sync"
                                         % (decision or "?", place), where))
            if row.get("Public Number", "") != numbers.get(record, ""):
                errors.append("%s: Public Number %s differs from chronology %s (run import)" % (
                    where, row.get("Public Number") or "-", numbers.get(record) or "-"))
            local_number, public_number = row.get("Local Number", ""), row.get("Public Number", "")
            if local_number != public_number:
                if not public_number:
                    kind = "Local Number without master Public Number: cleared at sync"
                elif not local_number:
                    kind = "Public Number not yet on Local: assigned at sync"
                else:
                    kind = "Local Number differs from Public Number: renumber at sync"
                warnings.append((kind, "%s [%s -> %s]" % (where, local_number or "-", public_number or "-")))
        by_id = {r.get("Record ID"): r for r in rows if r.get("Record ID")}
        for where, parent, relation in parents:
            if parent not in seen:
                errors.append("%s: Canonical / Parent Record ID %s not found in %s" % (where, parent, sheet))
            elif relation in SAME_ENTITY and by_id[parent].get("Publication Decision") == "ARCHIVE":
                errors.append("%s: canonical %s is itself ARCHIVE (point to the published card)" % (where, parent))
            elif relation in SAME_ENTITY and by_id[parent].get("Publication Decision") != "PUBLIC":
                warnings.append(("canonical card is not PUBLIC: alias leads nowhere yet", where))
        for record in sorted(parent_cycles(rows)):
            errors.append("%s: Canonical / Parent Record ID cycle through %s" % (sheet, record))
        claims = {}
        for row in rows:
            owner = canonical_target(row, by_id) or row.get("Record ID")
            for name in [row.get("Name", "")] + split_aliases(row.get("Aliases")):
                if name:
                    claims.setdefault(_alias_key(name), set()).add(owner)
        for row in rows:
            owner = canonical_target(row, by_id) or row.get("Record ID")
            others = claims.get(_alias_key(row.get("Name", "")), set()) - {owner}
            if row.get("Name") and others:
                warnings.append(("name matches another record or alias (possible duplicate / re-discovery; "
                                 "an ARCHIVE decision stands unless reopened)",
                                 "%s (%s) ~ %s" % (sheet, row.get("Record ID"), ", ".join(sorted(others)))))
            for alias in split_aliases(row.get("Aliases")):
                owners = claims.get(_alias_key(alias), set())
                if len(owners) > 1:
                    errors.append("%s (%s): alias %r is ambiguous, also claimed by %s" % (
                        sheet, row.get("Record ID"), alias, ", ".join(sorted(owners - {row.get("Record ID")}))))
        ids[sheet] = set(seen)
    for sheet, where, record, (target_sheet, target), relation, decision in cross_parents:
        target_rows = {r.get("Record ID"): r for r in workbook_rows.get(target_sheet, [])}
        target_row = target_rows.get(target)
        if target_sheet == sheet:
            errors.append("%s: use a plain Record ID for a canonical card on the same sheet" % where)
        elif target_row is None:
            errors.append("%s: Canonical / Parent Record ID %s:%s not found" % (where, target_sheet, target))
        elif relation in SAME_ENTITY and target_row.get("Publication Decision") != "PUBLIC":
            errors.append("%s: cross-sheet canonical %s:%s must be PUBLIC" % (where, target_sheet, target))
        elif cross_parent(target_row.get("Canonical / Parent Record ID", "")) == (sheet, record):
            errors.append("%s: Canonical / Parent Record ID cycle across sheets" % where)
        if relation and relation not in SAME_ENTITY:
            errors.append("%s: a cross-sheet canonical link must be a same-entity relation" % where)
        if decision != "ARCHIVE":
            errors.append("%s: a row linked to a canonical card on another sheet must be ARCHIVE" % where)
    for sheet in AUX:
        seen = set()
        for index, row in enumerate(workbook_rows.get(sheet, []), start=2):
            where = "%s row %d" % (sheet, index)
            if not row.get("Key"):
                errors.append("%s: empty Key" % where)
            elif row["Key"] in seen:
                errors.append("%s: duplicate Key %s" % (where, row["Key"]))
            seen.add(row.get("Key"))
            target = RECORD_TYPE_SHEET.get(row.get("Record Type"))
            if target is None:
                errors.append("%s: Record Type must be model or tool" % where)
            elif row.get("Record ID") not in ids.get(target, set()):
                errors.append("%s: Record ID %s not found in %s" % (where, row.get("Record ID"), target))
            if sheet == "Offers":
                # Empty Billing Unit means the standard unit label of the site
                # (e.g. input = USD per 1M input tokens); "other" needs its own.
                if row.get("Unit") not in OFFER_UNITS:
                    errors.append("%s (%s): Offer Unit %r is not a site unit" % (where, row.get("Key"), row.get("Unit")))
                elif row.get("Unit") == "other" and not row.get("Billing Unit"):
                    errors.append("%s (%s): Unit other requires Billing Unit" % (where, row.get("Key")))
    if snapshot is not None:
        for sheet in MAIN:
            master = {r.get("Record ID"): r for r in workbook_rows.get(sheet, [])}
            missing = sorted(set(snapshot[sheet]) - set(master))
            if missing:
                errors.append("%s: %d Local records absent from master: %s" % (
                    sheet, len(missing), ", ".join(missing[:10])))
            fields = records = 0
            for record_id, (local_row, _) in snapshot[sheet].items():
                row = master.get(record_id)
                if row is None:
                    continue
                diff = sum(1 for c in PUBLIC[sheet] if row.get(c, "") != local_row.get(c, ""))
                fields += diff
                records += bool(diff)
            drift[sheet] = {"records": records, "fields": fields}
    if aux is not None:
        for sheet in AUX:
            master = {r.get("Key"): r for r in workbook_rows.get(sheet, [])}
            missing = sorted(set(aux[sheet]) - set(master))
            natural = NATURAL_KEYS.get(sheet)
            if natural:
                present = {tuple(r.get(c, "") for c in natural) for r in master.values()}
                missing = [k for k in missing if tuple(aux[sheet][k].get(c, "") for c in natural) not in present]
            if missing:
                errors.append("%s: %d Local rows absent from master: %s" % (
                    sheet, len(missing), ", ".join(missing[:10])))
            fields = records = 0
            for key, values in aux[sheet].items():
                row = master.get(key)
                if row is None:
                    continue
                diff = sum(1 for c, v in values.items() if row.get(c, "") != v)
                fields += diff
                records += bool(diff)
            drift[sheet] = {"records": records, "fields": fields}
    return errors, warnings, drift


# ---------------------------------------------------------------- workbook IO

def read_workbook(path=WORKBOOK_PATH):
    """Return (rows per sheet, meta dict, extra headers per sheet); all text."""
    from openpyxl import load_workbook

    workbook = load_workbook(path)
    if "Meta" not in workbook.sheetnames:
        raise ValueError("%s is not a catalog master (no Meta sheet)" % path)
    meta = {_text(k): _text(v) for k, v in workbook["Meta"].iter_rows(min_row=2, values_only=True) if k}
    upgrading = meta.get("Schema") in UPGRADABLE_SCHEMAS
    if meta.get("Schema") != SCHEMA and not upgrading:
        raise ValueError("%s has schema %r, expected %r" % (path, meta.get("Schema"), SCHEMA))
    rows, extra = {}, {}
    for sheet, columns in [*MAIN.items(), *AUX.items(), ("Changelog", CHANGELOG)]:
        if sheet not in workbook.sheetnames:
            raise ValueError("%s lacks sheet %s" % (path, sheet))
        values = workbook[sheet].iter_rows(values_only=True)
        headers = [_text(h) for h in next(values, [])]
        absent = [c for c in columns if c not in headers]
        if upgrading and sheet in MAIN:
            absent = [c for c in absent if c not in (
                "Publication Decision", "Approx Precision", "Reason", "Canonical / Parent Record ID",
                "Decision Code", "Decision Date", "Decision Sources", "Relation Type", "Aliases")]
        if absent:
            raise ValueError("%s sheet %s lacks columns: %s" % (path, sheet, ", ".join(absent)))
        extra[sheet] = [h for h in headers if h and h not in columns]
        rows[sheet] = []
        for raw in values:
            row = {h: _text(v) for h, v in zip(headers, raw) if h}
            if any(row.values()):
                if upgrading and sheet in MAIN and not row.get("Publication Decision"):
                    row["Publication Decision"] = "PUBLIC" if row.get("Status") == "PUBLISHED" else "NEEDS_REVIEW"
                rows[sheet].append(row)
    if upgrading:
        meta["Schema"] = SCHEMA
    return rows, meta, extra


RULES_RU = [
    ("Назначение", "Единая каноническая база AIpediya: все Models и Tools — опубликованные, скрытые кандидаты, архив и новые находки. Запись сначала появляется и проверяется здесь; затем изменения переносятся в Local (проверенная синхронизация) и только по docs/RELEASE.md и отдельной команде владельца — в Production. Действующая книга одна: AI_CONTEXT/AIpediya_Model_Verification_Master.xlsx; версии v001–v012 — неактивные резервные копии."),
    ("Три разных понятия", "Publication Decision — редакционное решение (нужна ли самостоятельная карточка). Status — разрешено ли показывать запись в публичном каталоге при следующей проверенной синхронизации. On Local / On Production — фактическое наблюдаемое состояние контуров на дату проверки. Catalog Status — жизненный цикл продукта (active/deprecated/retired/archived), не редакционная ценность."),
    ("Status", "PUBLISHED — разрешено к публикации: Publication Decision=PUBLIC и нет блокирующих пробелов (Missing Data пусто; исключение — уже опубликованная на сайте запись, у которой не хватает только даты: она остаётся без Public Number до решения владельца). NEEDS_REVIEW — не показывать публично (ARCHIVE, NEEDS_REVIEW или PUBLIC с блокирующим пробелом). PUBLISHED не означает, что запись уже на сайте."),
    ("Publication Decision", "PUBLIC — подходит для самостоятельной публичной карточки: подтверждены идентичность, разработчик, назначение и тип выпуска, и понятно, почему нужна отдельная карточка. ARCHIVE — остаётся только во внутренней базе: доказанный дубль, alias или датированный идентификатор той же модели, режим/конфигурация, формат весов, технический checkpoint, семейный агрегат при наличии отдельных карточек, продукт не того листа, вне области каталога, отменённый выпуск, тестовый артефакт. NEEDS_REVIEW — есть конкретный нерешённый вопрос (в Reason). Закрытие продукта, малый размер, язык, страна или неизвестный бренд сами по себе не основание для ARCHIVE; недостаток сведений не маскировать ARCHIVE."),
    ("Decision Code", "Код решения. PUBLIC: CURRENT_RELEASE / HISTORICAL_RELEASE. ARCHIVE: DUPLICATE, ALIAS_SNAPSHOT, API_SNAPSHOT, CONFIGURATION, FORMAT_VARIANT, TECHNICAL_CHECKPOINT, FAMILY_AGGREGATE, MODEL_APP_SPLIT, OUT_OF_SCOPE, CANCELLED, TEST_ARTIFACT. NEEDS_REVIEW: IDENTITY, DATE, EXISTENCE, SCOPE, SOURCE."),
    ("Reason", "Индивидуальный комментарий: для ARCHIVE — установленный факт и отдельно редакционный вывод (не «мусор» и не «никому не нужно»); для NEEDS_REVIEW — точный незакрытый вопрос; для PUBLIC — зачем нужна отдельная карточка."),
    ("Decision Date / Decision Sources", "Дата редакционного решения (YYYY-MM-DD) и основания (URL через ' | '). Для ARCHIVE обязательны код, Reason, дата и источник (Decision Sources или Official Source)."),
    ("Canonical / Parent Record ID + Relation Type", "Связь с другой записью того же листа. DUPLICATE_OF / ALIAS_OF / FORMAT_OF / MODE_OF — та же сущность под другим именем или в другой форме: строка всегда ARCHIVE, канон — PUBLIC-запись, её имя и aliases ведут на канон. SNAPSHOT_OF / VARIANT_OF / MEMBER_OF — связь разных сущностей (версия, вариант, член семейства), не объединение. Цепочки к архиву и циклы запрещены. Цены и оценки между разными версиями не переносятся."),
    ("Aliases", "Альтернативные названия той же сущности через ' | ' (пробел, вертикальная черта, пробел). Все aliases ищутся; один alias не может принадлежать двум разным каноническим записям. Вариант или семейство — не alias."),
    ("Повторная находка", "Перед добавлением искать во всей книге (Name, Aliases, Developer, model ID, Official Source), включая ARCHIVE. Повторная находка не отменяет ARCHIVE: пересмотр — только явно, с новым основанием и записью в Changelog."),
    ("Approx Precision", "Вычисляется из Approx Date: year / month / day."),
    ("Record ID", "Slug записи. Стабильная идентичность: после создания не меняется и не переиспользуется (URL /models/<id>, /tools/<id>), не зависит от имени, сортировки, публикации или объединения."),
    ("Public Number", "Вычисляется, вручную не править. Хронологический номер по verified release date среди PUBLISHED (отдельно Models и Tools): точная дата, иначе ≈ дата (начало периода); при равной дате — по имени, затем Record ID. Историческая вставка сдвигает последующие номера — каждое изменение пишется в Changelog. Непубличные записи и записи без подтверждённой даты номера не получают (сортируются в конце). Local получает номера через sync-local (правило владельца 2026-09-26), прежние номера сохраняются в ревизиях Local; Production — только отдельным выпуском."),
    ("Даты", "Exact Release Date — только подтверждённая (YYYY-MM-DD). Иначе надёжная дата первого публичного существования — Approx Date: ≈YYYY-MM или ≈YYYY-MM-DD. Обе сразу не заполнять. Различать анонс, preview, фактический доступ, публикацию весов и обновление страницы. Дату git-коммита за дату выпуска не выдавать. Фиктивную дату ради номера не ставить."),
    ("Проверка", "Official Source — первичный источник; Secondary Source — независимое подтверждение (качество и ограничения — в Notes); Last Verified — дата реальной проверки. У PUBLISHED обязательны Name, Developer, Official Source, Last Verified. Ссылки импорта в Import Sources (unverified) ничего не подтверждают."),
    ("Missing Data", "Вычисляется: чего не хватает (publication_decision, name, developer, category, date, official_source, last_verified, description_en; для ARCHIVE — decision_code, reason, decision_date, decision_sources, last_verified, canonical_record для дублей). Неприменимые поля (контекст или benchmark у инструмента) пробелами не считаются. Пусто = формально достаточно, не «всё подтверждено»."),
    ("On Local / On Production", "Факт публичности записи в Local (база) и на Production (публичный sitemap, только чтение) на дату из Meta. Расхождение со Status = работа для проверенной синхронизации."),
    ("Local Number", "Номер, который сейчас стоит в Local. Отличие от Public Number = план перенумерации, применяемый только отдельным решением."),
    ("Связанные листы", "Offers, Evaluations, Access, Facts, Origins, Tool Platforms — цены, оценки, доступ, факты, страны и платформы; связь по Record Type + Record ID. Key — стабильный ключ строки. Offers: цена в USD; пустой Billing Unit означает стандартную единицу сайта (input = за 1M входных токенов и т. д.), Unit=other требует Billing Unit. Evaluations: Public=YES только при подтверждённой идентичности и праве показа; результаты разработчика не выдавать за независимые. Facts с ключами проверки (scoped_verification и т. п.) — внутренние, на сайт не переносятся."),
    ("Локализации", "EN и RU — редактируемые колонки. Остальные 20 локалей лежат в *Other Locales (JSON) и выводятся из английского конвейером переводов."),
    ("Неизвестное", "Оставлять пустым, не нулём и не догадкой."),
    ("Changelog", "Журнал изменений master (только дописывается): импорт, решения, перенумерация, правки полей с Before/After и причиной."),
    ("Пополнение", "Поиск во всей базе → обновление существующей записи или новая запись с новым Record ID → проверка источников → редакционное решение → refresh + check → пробная синхронизация на изолированной копии Local → приёмка владельцем → отдельно разрешённая публикация."),
    ("Команды", "Пересчёт производных полей и Meta: .\\.venv\\Scripts\\python.exe manage.py catalog_master refresh | Импорт новых записей из Local (master не перетирается): … catalog_master import [--production] | Проверка: … catalog_master check [--production] | План синхронизации master → Local: … catalog_master sync-local [--apply] (без --apply ничего не пишет; применять сначала к копии через AIPEDIA_DB); после --apply выполнить import, чтобы книга получила новые ключи и фактическое On Local. Что переносится — docs/CATALOG_MASTER.md."),
    ("Production", "Таблица и команды не публикуют и не меняют Production; --production только читает публичный sitemap."),
]


def write_workbook(workbook_rows, meta, path=WORKBOOK_PATH, extra=None):
    """Write the whole master atomically (temp file, then replace)."""
    from openpyxl import Workbook
    from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
    from openpyxl.formatting.rule import FormulaRule
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation

    extra = extra or {}

    def fill(color):
        # Excel reads conditional-format fills from bgColor, not fgColor.
        return PatternFill("solid", fgColor=color, bgColor=color)

    header_fills = {"workflow": "1F3A5F", "verification": "7A4B00", "public": "2F5D3A", "service": "6B7280"}
    workbook = Workbook()
    workbook.remove(workbook.active)

    def write_sheet(name, columns, rows, groups=None, widths=None):
        sheet = workbook.create_sheet(name)
        headers = columns + [h for h in extra.get(name, []) if h not in columns]
        sheet.append(headers)
        for row in rows:
            values = []
            for header in headers:
                value = row.get(header, "")
                if isinstance(value, str):
                    value = ILLEGAL_CHARACTERS_RE.sub("", value)
                values.append(value or None)
            sheet.append(values)
            for cell in sheet[sheet.max_row]:
                if isinstance(cell.value, str) and cell.value.startswith("="):
                    cell.data_type = "s"
        for index, header in enumerate(headers, start=1):
            cell = sheet.cell(row=1, column=index)
            group = (groups or {}).get(header, "service")
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor=header_fills[group])
            cell.alignment = Alignment(wrap_text=True, vertical="center")
            letter = get_column_letter(index)
            sheet.column_dimensions[letter].width = (widths or {}).get(header, 16)
            sheet.column_dimensions[letter].number_format = "@"
        sheet.row_dimensions[1].height = 32
        sheet.auto_filter.ref = "A1:%s%d" % (get_column_letter(len(headers)), max(len(rows) + 1, 2))
        return sheet, {h: get_column_letter(i) for i, h in enumerate(headers, start=1)}

    def validation(sheet, letter, values):
        rule = DataValidation(type="list", formula1='"%s"' % ",".join(values), allow_blank=True)
        rule.error = "Allowed: " + ", ".join(values)
        rule.showErrorMessage = True
        sheet.add_data_validation(rule)
        rule.add("%s2:%s%d" % (letter, letter, MAX_VALIDATED_ROW))

    widths = {
        "Record ID": 34, "Status": 15, "Public Number": 9, "Name": 34, "Developer": 24,
        "Exact Release Date": 13, "Approx Date": 12, "Missing Data": 40,
        "Official Source": 40, "Secondary Source": 32, "Last Verified": 12,
        "On Local": 9, "On Production": 11, "Notes": 50, "Description EN": 50,
        "Description RU": 50, "Source URL": 36, "Official URL": 32,
    }
    last = MAX_VALIDATED_ROW
    for name, columns in MAIN.items():
        groups = {c: "workflow" for c in WORKFLOW}
        groups.update({c: "verification" for c in VERIFICATION})
        groups.update({c: "public" for c in (MODEL_PUBLIC if name == "Models" else TOOL_PUBLIC)})
        sheet, col = write_sheet(name, columns, workbook_rows[name], groups, widths)
        sheet.freeze_panes = "F2"
        validation(sheet, col["Status"], STATUSES)
        validation(sheet, col["Publication Decision"], DECISIONS)
        validation(sheet, col["Decision Code"], ALL_DECISION_CODES)
        validation(sheet, col["Relation Type"], RELATION_TYPES)
        decision = col["Publication Decision"]
        for value, color in (("PUBLIC", "C6E0B4"), ("ARCHIVE", "D9D9D9"), ("NEEDS_REVIEW", "FFF4CE")):
            sheet.conditional_formatting.add(
                "%s2:%s%d" % (decision, decision, last),
                FormulaRule(formula=['$%s2="%s"' % (decision, value)], fill=fill(color)))
        validation(sheet, col["On Local"], YES_NO)
        validation(sheet, col["On Production"], YES_NO)
        validation(sheet, col["Category"], MODEL_CATEGORIES if name == "Models" else TOOL_CATEGORIES)
        validation(sheet, col["Catalog Status"], CATALOG_STATUSES)
        if name == "Models":
            validation(sheet, col["Release Stage"], STAGES)
            validation(sheet, col["Open Weights"], YES_NO)
        else:
            validation(sheet, col["Local Execution"], LOCAL_EXECUTION)
        status = col["Status"]
        for value, color in (("PUBLISHED", "C6E0B4"), ("NEEDS_REVIEW", "FFF4CE")):
            sheet.conditional_formatting.add(
                "%s2:%s%d" % (status, status, last),
                FormulaRule(formula=['$%s2="%s"' % (status, value)], fill=fill(color)))
        missing = col["Missing Data"]
        sheet.conditional_formatting.add(
            "%s2:%s%d" % (missing, missing, last),
            FormulaRule(formula=['$%s2<>""' % missing], fill=fill("F8CBAD")))
        for column in ("On Local", "On Production"):
            letter = col[column]
            sheet.conditional_formatting.add(
                "%s2:%s%d" % (letter, letter, last),
                FormulaRule(formula=['OR(AND($%s2="NEEDS_REVIEW",$%s2="YES"),AND($%s2="PUBLISHED",$%s2="NO"))'
                                     % (status, letter, status, letter)], fill=fill("F4B6B6")))
    for name, columns in AUX.items():
        sheet, col = write_sheet(name, columns, workbook_rows[name], widths={"Key": 18, "Record ID": 34})
        sheet.freeze_panes = "D2"
        validation(sheet, col["Record Type"], ["model", "tool"])
    sheet, _ = write_sheet("Changelog", CHANGELOG, workbook_rows.get("Changelog", []),
                           widths={"Timestamp (UTC)": 21, "Record ID": 34, "Reason": 60})
    sheet.freeze_panes = "A2"

    meta_sheet = workbook.create_sheet("Meta")
    meta_sheet.append(["Key", "Value"])
    for key, value in meta.items():
        meta_sheet.append([key, value])
    meta_sheet.column_dimensions["A"].width = 28
    meta_sheet.column_dimensions["B"].width = 70

    rules = workbook.create_sheet("Rules")
    rules.append(["Правило", "Описание"])
    for item in RULES_RU:
        rules.append(list(item))
    rules.column_dimensions["A"].width = 26
    rules.column_dimensions["B"].width = 130
    for (cell,) in rules.iter_rows(min_col=2, max_col=2):
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    for (cell,) in rules.iter_rows(min_col=1, max_col=1):
        cell.font = Font(bold=True)

    lists = workbook.create_sheet("Lists")
    for index, (title, values) in enumerate([
            ("Status", STATUSES), ("Publication Decision", DECISIONS), ("YES/NO", YES_NO),
            ("Decision Code", ALL_DECISION_CODES), ("Relation Type", RELATION_TYPES),
            ("Model Category", MODEL_CATEGORIES),
            ("Tool Category", TOOL_CATEGORIES), ("Catalog Status", CATALOG_STATUSES),
            ("Release Stage", STAGES), ("Local Execution", LOCAL_EXECUTION)], start=1):
        lists.cell(row=1, column=index, value=title).font = Font(bold=True)
        for offset, value in enumerate(values, start=2):
            lists.cell(row=offset, column=index, value=value)
        lists.column_dimensions[get_column_letter(index)].width = 20

    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    handle, temp = tempfile.mkstemp(suffix=".xlsx", dir=path.parent)
    os.close(handle)
    try:
        workbook.save(temp)
        os.replace(temp, path)
    except PermissionError as exc:
        raise PermissionError("%s is locked (close it in Excel and retry)" % path) from exc
    finally:
        if os.path.exists(temp):
            os.remove(temp)
